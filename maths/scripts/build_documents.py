#!/usr/bin/env python3
"""Build an A4 maths paper with answers, or a mistake summary, from JSON.

The caller supplies verified mathematical content. This script validates the
schema and scoring, draws exact-coordinate PNGs and lays out editable Word math.
It does not OCR, invent evidence, diagnose a learner, or certify mathematics.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

from draw_figures import draw_figure


def math_nodes(value):
    """A deliberately small native OMML AST: text/list/frac/sup/sub/sqrt."""
    if isinstance(value, (str, int, float)):
        run = OxmlElement("m:r")
        props = OxmlElement("m:rPr")
        sty = OxmlElement("m:sty")
        sty.set(qn("m:val"), "p")
        props.append(sty)
        run.append(props)
        text = OxmlElement("m:t")
        text.set(qn("xml:space"), "preserve")
        text.text = str(value)
        run.append(text)
        return [run]
    if isinstance(value, list):
        return [n for part in value for n in math_nodes(part)]
    if not isinstance(value, dict) or len(value) != 1:
        raise ValueError(f"Invalid mathematical expression: {value!r}")
    key, val = next(iter(value.items()))
    tags = {"frac": ("m:f", ["m:num", "m:den"]),
            "sup": ("m:sSup", ["m:e", "m:sup"]),
            "sub": ("m:sSub", ["m:e", "m:sub"])}
    if key == "sqrt":
        root = OxmlElement("m:rad")
        props = OxmlElement("m:radPr")
        hide = OxmlElement("m:degHide")
        hide.set(qn("m:val"), "1")
        props.append(hide)
        root.append(props)
        root.append(OxmlElement("m:deg"))
        expr = OxmlElement("m:e")
        expr.extend(math_nodes(val))
        root.append(expr)
        return [root]
    if key not in tags or not isinstance(val, list) or len(val) != 2:
        raise ValueError(f"Unsupported mathematical expression: {value!r}")
    name, children = tags[key]
    root = OxmlElement(name)
    for tag, child in zip(children, val):
        element = OxmlElement(tag)
        element.extend(math_nodes(child))
        root.append(element)
    return [root]


def rich(p, content, bold=False):
    if isinstance(content, str):
        p.add_run(content).bold = bold
        return p
    if not isinstance(content, list):
        raise ValueError("Paragraph must be a string or list of text/math parts")
    for part in content:
        if isinstance(part, str):
            p.add_run(part).bold = bold
        elif isinstance(part, dict) and set(part) == {"math"}:
            omath = OxmlElement("m:oMath")
            omath.extend(math_nodes(part["math"]))
            p._p.append(omath)
        else:
            raise ValueError(f"Invalid rich text part: {part!r}")
    return p


def validate(spec, mode="paper"):
    if not spec.get("title"):
        raise ValueError("title is required")
    questions = ([q for section in spec["sections"] for q in section["questions"]]
                 if spec.get("sections") else spec.get("questions", []))
    if not questions:
        raise ValueError("nonempty sections or questions are required")
    for q in questions:
        for field in ("stem", "solution"):
            if not isinstance(q.get(field), list) or not q[field]:
                raise ValueError(f"Question {q.get('id')} {field} must be a nonempty paragraph array")
            for content in q[field]:
                if not isinstance(content, (str, list)):
                    raise ValueError(f"Question {q.get('id')} has invalid {field} paragraph")
    ids = [q["id"] for q in questions]
    if mode == "mistakes":
        if len(ids) != len(set(ids)):
            raise ValueError("Question ids must be unique")
        records = spec.get("mistake_records", [])
        if not records:
            raise ValueError("No mistake records; do not fabricate a summary")
        for record in records:
            if record.get("question_id") not in ids or not record.get("record_id") or not record.get("evidence"):
                raise ValueError("Each mistake record needs a question, record_id and evidence")
        for q in questions:
            q.setdefault("points", 0)
            if not q.get("stem") or not q.get("solution"):
                raise ValueError("A mistake question requires a stem and verified solution")
        return questions
    if ids != list(range(1, len(ids)+1)):
        raise ValueError("Question ids must be unique consecutive integers starting at 1")
    if sum(q["points"] for q in questions) != spec["total_points"]:
        raise ValueError("Question points do not equal total_points")
    if spec.get("question_count", len(ids)) != len(ids):
        raise ValueError("Question count mismatch")
    for q in questions:
        if not isinstance(q["points"], int) or q["points"] <= 0:
            raise ValueError(f"Question {q['id']} must have positive integer points")
        if not q.get("stem") or not q.get("answer") or not q.get("solution"):
            raise ValueError(f"Question {q['id']} requires stem, answer and solution")
        if not q.get("source") or q["source"].get("kind") not in (
                "original", "variant", "user_upload", "verified_past_paper"):
            raise ValueError(f"Question {q['id']} needs a source kind")
        if q["source"]["kind"] == "verified_past_paper" and not q["source"].get("reference"):
            raise ValueError("A verified past-paper question needs an auditable reference")
        if "options" in q and len(q["options"]) != 4:
            raise ValueError("Choice questions must have four options")
        if sum(item["points"] for item in q.get("scoring", [])) != q["points"]:
            raise ValueError(f"Question {q['id']} scoring does not sum to question points")
        if any(item["points"] <= 0 or not item.get("text") for item in q["scoring"]):
            raise ValueError("Every scoring step needs positive points and a description")
    return questions


def set_font(style, font, size, bold=False):
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for attr in list(fonts.attrib):
        if "theme" in attr.lower():
            del fonts.attrib[attr]
    fonts.set(qn("w:eastAsia"), font)


def base_document(spec, title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(17), Mm(17)
    sec.left_margin, sec.right_margin = Mm(20), Mm(20)
    sec.header_distance, sec.footer_distance = Mm(8), Mm(8)
    for grid in sec._sectPr.findall(qn("w:docGrid")):
        sec._sectPr.remove(grid)
    default_cjk = "Songti SC" if sys.platform == "darwin" else ("SimSun" if sys.platform == "win32" else "Noto Serif CJK SC")
    cjk_font = os.environ.get("MATHS_CJK_FONT", default_cjk)
    set_font(doc.styles["Normal"], cjk_font, 11)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.0
    normal.space_after = Pt(4)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    doc.styles["Normal"].element.get_or_add_pPr().append(snap)
    set_font(doc.styles["Title"], cjk_font, 19, True)
    set_font(doc.styles["Heading 1"], cjk_font, 13, True)
    set_font(doc.styles["Heading 2"], cjk_font, 11, True)
    for name in ("Title", "Heading 1", "Heading 2"):
        for border in doc.styles[name].element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
        doc.styles[name].paragraph_format.space_before = Pt(7)
        doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.core_properties.author = "maths"
    doc.core_properties.title = title
    doc.core_properties.subject = spec.get("topic", "初中数学")
    # Only page numbers are included: they help keep printed work in order.
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    p.add_run(" 页")
    p = doc.add_paragraph(title, "Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def paragraph(doc, content, **kwargs):
    p = doc.add_paragraph()
    rich(p, content)
    for key, value in kwargs.items():
        setattr(p.paragraph_format, key, value)
    return p


def question_block(doc, q, figures, stem_only=False):
    pars = []
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = bool(q.get("paper_page_before")) and not stem_only
    p.add_run(f"{q['id']}．" + (f"（{q['points']} 分）" if q.get("points") else "")).bold = True
    rich(p, q["stem"][0])
    pars.append(p)
    for content in q["stem"][1:]:
        pars.append(paragraph(doc, content))
    if q.get("options"):
        columns = q.get("options_columns", 2)
        if columns not in (2, 4):
            raise ValueError("options_columns must be 2 or 4")
        for i in range(0, 4, columns):
            p = doc.add_paragraph()
            for n in range(1, columns):
                p.paragraph_format.tab_stops.add_tab_stop(Mm(164*n/columns))
            for j in range(i, i+columns):
                p.add_run(f"{'ABCD'[j]}．")
                rich(p, q["options"][j])
                if j < i+columns-1:
                    p.add_run("\t")
            pars.append(p)
    if q["id"] in figures:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        shape = run.add_picture(str(figures[q["id"]]), width=Mm(q.get("figure_width_mm", 72)))
        shape._inline.docPr.set("descr", q.get("figure_alt", f"第{q['id']}题数学配图"))
        p.paragraph_format.line_spacing = Pt(shape.height / 12700 + 6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pars.append(p)
    # Keep a stem, its options and its figure on the same page. Answer room can split.
    for p in pars[:-1]:
        p.paragraph_format.keep_with_next = True
    for p in pars:
        p.paragraph_format.keep_together = True
    if not stem_only:
        lines = q.get("space_lines", 0)
        for _ in range(lines):
            p = doc.add_paragraph(" ")
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(18)


def answer_block(doc, q):
    start = len(doc.paragraphs)
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.page_break_before = bool(q.get("answer_page_before"))
    p.add_run(f"第 {q['id']} 题  {q['points']} 分")
    if len(q["scoring"]) == 1:
        p.add_run("　答案：")
        rich(p, q["answer"])
        for content in q["solution"]:
            paragraph(doc, content, keep_together=True)
        for p in doc.paragraphs[start:-1]:
            p.paragraph_format.keep_with_next = True
        return
    paragraph(doc, ["答案："] + (q["answer"] if isinstance(q["answer"], list) else [q["answer"]]),
              keep_with_next=True)
    for content in q["solution"]:
        paragraph(doc, content, keep_together=True)
    for item in q["scoring"]:
        paragraph(doc, f"得分点 {item['points']} 分：{item['text']}", keep_together=True,
                  space_after=Pt(3))
    for p in doc.paragraphs[start:-1]:
        p.paragraph_format.keep_with_next = True


def prepare_figures(questions, out_dir, input_dir):
    paths, manifest = {}, []
    for q in questions:
        if not q.get("figure"):
            continue
        target = out_dir / "figures" / f"q{q['id']:02}.png"
        fig = q["figure"]
        if "path" in fig:
            source = (input_dir / fig["path"]).resolve()
            with Image.open(source) as im:
                if im.format != "PNG":
                    raise ValueError("External figure must be a PNG")
                size = list(im.size)
            target.parent.mkdir(parents=True, exist_ok=True)
            if source != target.resolve():
                shutil.copy2(source, target)
            info = {"path": str(target), "pixels": size, "method": "supplied_png"}
        else:
            info = draw_figure(fig, target)
        paths[q["id"]] = target
        manifest.append({"question_id": q["id"], **info})
    return paths, manifest


def build_paper(spec, out_dir, figures):
    doc = base_document(spec, spec["title"])
    p = paragraph(doc, f"{spec['minutes']} 分钟  满分 {spec['total_points']} 分  共 {spec['question_count']} 题")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "姓名：________________  日期：________________  得分：________")
    paragraph(doc, spec.get("notice", "独立完成题目，再查看文后的答案与解析。"))
    paragraph(doc, "请在题后空白处写出必要步骤。答案与详细解析附在试卷后，请完成后再查看。")
    for section in spec["sections"]:
        points = sum(q["points"] for q in section["questions"])
        heading = doc.add_heading(f"{section['title']}  共 {points} 分", 1)
        heading.paragraph_format.page_break_before = bool(section.get("page_before"))
        for q in section["questions"]:
            question_block(doc, q, figures)
    doc.add_heading("答案与详细解析", 1).paragraph_format.page_break_before = True
    paragraph(doc, "解答题按列出的关键步骤给分；等价的正确解法同样得分。每题得分不超过题目分值。")
    for section in spec["sections"]:
        for q in section["questions"]:
            answer_block(doc, q)
    path = out_dir / f"{safe_name(spec['title'])}.docx"
    doc.save(path)
    return path


def build_mistakes(spec, out_dir, figures, questions):
    if not spec.get("mistake_records"):
        raise ValueError("No mistake_records: do not fabricate a personal mistake summary")
    title = spec.get("mistakes_title", f"{spec.get('topic', '数学')} 错题整理")
    doc = base_document(spec, title)
    paragraph(doc, spec.get("mistakes_notice", "依据已保存的错题和作答记录整理。"))
    by_id = {q["id"]: q for q in questions}
    for i, record in enumerate(spec["mistake_records"]):
        q = by_id[record["question_id"]]
        paragraph(doc, f"记录标识：{record['record_id']}　知识点：{record['topic']}", page_break_before=bool(i))
        paragraph(doc, f"依据：{record['evidence']}")
        doc.add_heading("原题", 1)
        question_block(doc, q, figures, stem_only=True)
        for heading, key in [("易错点", "pitfalls"), ("出错原因", "cause"), ("改善方法", "improvement")]:
            doc.add_heading(heading, 1)
            values = record.get(key)
            if not values:
                values = ["尚无足够的实际作答证据，不能判断。"]
            for content in values:
                paragraph(doc, content)
        doc.add_heading("规范解法", 1)
        for content in q["solution"]:
            paragraph(doc, content, keep_together=True)
    path = out_dir / f"{safe_name(title)}.docx"
    doc.save(path)
    return path


def safe_name(title):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip(" .") or "maths"


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--out-dir", required=True, type=Path)
    ap.add_argument("--mode", choices=("paper", "mistakes"), default="paper")
    ap.add_argument("--overwrite", action="store_true", help="Explicitly rebuild an existing run directory for QA; do not use for a new historical export")
    args = ap.parse_args()
    spec = json.loads(args.input.read_text(encoding="utf-8"))
    questions = validate(spec, args.mode)
    if args.out_dir.exists() and any(args.out_dir.iterdir()) and not args.overwrite:
        raise ValueError("Output directory is not empty. Use a new run directory; --overwrite is only for an intentional rebuild.")
    args.out_dir.mkdir(parents=True, exist_ok=True)
    used = questions
    if args.mode == "mistakes":
        ids = {r["question_id"] for r in spec["mistake_records"]}
        used = [q for q in questions if q["id"] in ids]
    figures, figure_manifest = prepare_figures(used, args.out_dir, args.input.parent)
    if args.mode == "paper":
        output = build_paper(spec, args.out_dir, figures)
    else:
        output = build_mistakes(spec, args.out_dir, figures, questions)
    manifest = {"mode": args.mode, "document": str(output.resolve()),
                "question_count": len(used), "total_points": spec.get("total_points") if args.mode == "paper" else None,
                "figures": figure_manifest, "layout_qa": "pending_render_and_visual_review",
                "content_qa": "caller_must_verify_mathematics_and_evidence"}
    (args.out_dir / f"manifest-{args.mode}.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
